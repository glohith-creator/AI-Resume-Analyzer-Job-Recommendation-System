import io
import re
import pandas as pd

from pypdf import PdfReader
from docx import Document

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


# =========================================================
# 1. RESUME TEXT EXTRACTION
# =========================================================

def extract_pdf_text(file_bytes):
    """
    Extract text from all pages of a PDF resume.
    """

    pdf_file = io.BytesIO(file_bytes)
    reader = PdfReader(pdf_file)

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n".join(pages)


def extract_docx_text(file_bytes):
    """
    Extract text from a DOCX resume.
    """

    docx_file = io.BytesIO(file_bytes)
    document = Document(docx_file)

    text_parts = []

    # Normal paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text)

    return "\n".join(text_parts)


def extract_resume_text(uploaded_file):
    """
    Automatically detect PDF or DOCX and extract text.
    """

    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        return extract_pdf_text(file_bytes)

    elif filename.endswith(".docx"):
        return extract_docx_text(file_bytes)

    else:
        raise ValueError(
            "Unsupported file format. Please upload a PDF or DOCX file."
        )


# =========================================================
# 2. TEXT CLEANING
# =========================================================

def clean_text(text):
    """
    Clean resume text for NLP processing.
    """

    text = text.lower()

    # Replace line breaks with spaces
    text = re.sub(r"[\r\n\t]+", " ", text)

    # Remove excessive spaces
    text = re.sub(r"\s+", " ", text)

    return text.strip()


# =========================================================
# 3. LOAD SKILL DICTIONARY
# =========================================================

def load_skill_dictionary(csv_path):
    """
    Load skills and aliases from skill_dictionary.csv.

    Expected columns:

    skill,aliases

    Example:

    Python,"python,py"
    JavaScript,"javascript,js"
    Machine Learning,"machine learning,ml"
    """

    df = pd.read_csv(csv_path)

    required_columns = {"skill", "aliases"}

    if not required_columns.issubset(df.columns):
        raise ValueError(
            "skill_dictionary.csv must contain 'skill' and 'aliases' columns."
        )

    skill_dictionary = {}

    for _, row in df.iterrows():

        skill = str(row["skill"]).strip()

        aliases = str(row["aliases"]).split(",")

        terms = [skill]

        for alias in aliases:
            alias = alias.strip()

            if alias:
                terms.append(alias)

        skill_dictionary[skill] = terms

    return skill_dictionary


# =========================================================
# 4. SKILL EXTRACTION
# =========================================================

def extract_skills(resume_text, skill_dictionary):
    """
    Detect known technical skills from the resume.
    """

    cleaned_resume = clean_text(resume_text)

    detected_skills = []

    for skill, aliases in skill_dictionary.items():

        found = False

        for alias in aliases:

            alias = clean_text(alias)

            if not alias:
                continue

            # Handle symbols such as C++, C#, .NET
            escaped_alias = re.escape(alias)

            # Normal word boundary matching
            pattern = r"(?<!\w)" + escaped_alias + r"(?!\w)"

            if re.search(pattern, cleaned_resume):
                found = True
                break

        if found:
            detected_skills.append(skill)

    return sorted(set(detected_skills))


# =========================================================
# 5. LOAD JOB ROLE DATA
# =========================================================

def load_job_roles(csv_path):
    """
    Load job-role information from job_roles.csv.

    Expected columns:

    role,description,skills,roadmap

    Skills and roadmap items should be separated by '|'.
    """

    df = pd.read_csv(csv_path)

    required_columns = {
        "role",
        "description",
        "skills",
        "roadmap"
    }

    if not required_columns.issubset(df.columns):
        raise ValueError(
            "job_roles.csv must contain: "
            "role, description, skills, roadmap"
        )

    roles = []

    for _, row in df.iterrows():

        skills = [
            skill.strip()
            for skill in str(row["skills"]).split("|")
            if skill.strip()
        ]

        roadmap = [
            step.strip()
            for step in str(row["roadmap"]).split("|")
            if step.strip()
        ]

        roles.append({
            "role": str(row["role"]).strip(),
            "description": str(row["description"]).strip(),
            "skills": skills,
            "roadmap": roadmap
        })

    return roles


# =========================================================
# 6. TF-IDF SIMILARITY
# =========================================================

def calculate_similarity(resume_text, role_description):
    """
    Calculate TF-IDF cosine similarity between
    resume text and a job-role description.
    """

    resume = clean_text(resume_text)
    role = clean_text(role_description)

    documents = [
        resume,
        role
    ]

    vectorizer = TfidfVectorizer(
        stop_words="english",
        ngram_range=(1, 2)
    )

    matrix = vectorizer.fit_transform(documents)

    similarity = cosine_similarity(
        matrix[0:1],
        matrix[1:2]
    )[0][0]

    return float(similarity)


# =========================================================
# 7. SKILL MATCH SCORE
# =========================================================

def calculate_skill_score(resume_skills, required_skills):
    """
    Calculate percentage of required skills
    found in the resume.
    """

    resume_set = {
        skill.lower().strip()
        for skill in resume_skills
    }

    required_set = {
        skill.lower().strip()
        for skill in required_skills
    }

    if not required_set:
        return 0.0

    matched = resume_set.intersection(required_set)

    return len(matched) / len(required_set)


# =========================================================
# 8. FINAL ROLE MATCH SCORE
# =========================================================

def calculate_match_score(
    resume_text,
    resume_skills,
    role
):
    """
    Final score combines:

    60% TF-IDF similarity
    40% skill matching
    """

    tfidf_score = calculate_similarity(
        resume_text,
        role["description"]
    )

    skill_score = calculate_skill_score(
        resume_skills,
        role["skills"]
    )

    final_score = (
        (tfidf_score * 0.60)
        +
        (skill_score * 0.40)
    )

    return round(final_score * 100, 2)


# =========================================================
# 9. ANALYZE ALL JOB ROLES
# =========================================================

def analyze_resume(
    resume_text,
    resume_skills,
    job_roles
):
    """
    Compare the resume against every job role.
    """

    results = []

    for role in job_roles:

        score = calculate_match_score(
            resume_text,
            resume_skills,
            role
        )

        resume_skill_set = {
            skill.lower().strip()
            for skill in resume_skills
        }

        role_skill_set = {
            skill.lower().strip()
            for skill in role["skills"]
        }

        matched_skills = [
            skill
            for skill in role["skills"]
            if skill.lower().strip() in resume_skill_set
        ]

        missing_skills = [
            skill
            for skill in role["skills"]
            if skill.lower().strip() not in resume_skill_set
        ]

        results.append({
            "role": role["role"],
            "score": score,
            "matched_skills": matched_skills,
            "missing_skills": missing_skills,
            "roadmap": role["roadmap"]
        })

    # Highest score first
    results.sort(
        key=lambda item: item["score"],
        reverse=True
    )

    return results


# =========================================================
# 10. TOP 3 RECOMMENDATIONS
# =========================================================

def get_top_roles(results, number=3):
    """
    Return the top N recommended job roles.
    """

    return results[:number]


# =========================================================
# 11. GET SELECTED ROLE
# =========================================================

def get_role_result(results, role_name):
    """
    Find the analysis result for a selected role.
    """

    for result in results:

        if result["role"] == role_name:
            return result

    return None


# =========================================================
# 12. GENERATE LEARNING ROADMAP
# =========================================================

def generate_roadmap(role_result):
    """
    Generate a simple roadmap based on missing skills.
    """

    if role_result is None:
        return []

    roadmap = []

    missing_skills = role_result["missing_skills"]

    # First focus on missing skills
    for skill in missing_skills:
        roadmap.append(
            f"Learn and practice {skill}"
        )

    # Add predefined role roadmap
    for step in role_result["roadmap"]:

        if step not in roadmap:
            roadmap.append(step)

    return roadmap


# =========================================================
# 13. COMPLETE ANALYSIS PIPELINE
# =========================================================

def process_resume(
    uploaded_file,
    skill_dictionary,
    job_roles
):
    """
    Complete pipeline:

    Upload
       ↓
    Extract text
       ↓
    Clean text
       ↓
    Extract skills
       ↓
    Match job roles
       ↓
    Return results
    """

    resume_text = extract_resume_text(
        uploaded_file
    )

    if not resume_text.strip():
        raise ValueError(
            "No readable text was found in the resume."
        )

    cleaned_text = clean_text(
        resume_text
    )

    detected_skills = extract_skills(
        cleaned_text,
        skill_dictionary
    )

    results = analyze_resume(
        cleaned_text,
        detected_skills,
        job_roles
    )

    top_roles = get_top_roles(
        results,
        number=3
    )

    return {
        "resume_text": resume_text,
        "cleaned_text": cleaned_text,
        "detected_skills": detected_skills,
        "results": results,
        "top_roles": top_roles
    }